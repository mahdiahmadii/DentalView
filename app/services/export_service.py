"""
app.services.export_service
============================
Pure data-export helpers (CSV / Excel) with no Qt dependency, so they can
be tested and reused (e.g. from a future CLI or scheduled export) without
pulling in the UI stack.
"""

import csv


def export_detections_csv(path, rows):
    """rows: iterable of (class_name, confidence_str, box_str) tuples."""
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Class", "Confidence", "Box (x1,y1,x2,y2)"])
        for row in rows:
            writer.writerow(row)


def export_patients_excel(path, patients, problem_counts: dict):
    """patients: list of patient dicts (from Database.search_patients).
    problem_counts: {patient_id: int} pre-computed via
    CaseService.count_patient_problems.

    Raises ImportError if openpyxl is not installed — the caller is
    expected to catch this and show an actionable message.
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Patients"

    headers = ["Full Name", "Patient Code", "Sex", "Phone Number", "Teeth With Problems"]
    ws.append(headers)
    header_fill = PatternFill(start_color="4F7CFF", end_color="4F7CFF", fill_type="solid")
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for patient in patients:
        ws.append([
            patient.get("full_name") or "",
            patient.get("patient_code") or "",
            patient.get("sex") or "",
            patient.get("phone") or "",
            problem_counts.get(patient["id"], 0),
        ])

    for col_idx, header in enumerate(headers, start=1):
        max_len = len(header)
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            max_len = max(max_len, len(str(value)) if value is not None else 0)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max_len + 4

    wb.save(path)
